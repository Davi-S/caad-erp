import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_PATH = PROJECT_ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))


@pytest.fixture
def make_workbook():
    """Return a factory that creates a workbook with a single configured sheet."""

    def _build(sheet_name: str, headers: list[str]) -> Workbook:
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
        return wb

    return _build


@pytest.fixture
def products_workbook(make_workbook) -> Workbook:
    """Create an in-memory workbook with a canonical Products sheet header."""

    return make_workbook(
        "Products",
        ["ProductID", "ProductName", "SellPrice", "IsActive"],
    )


@pytest.fixture
def salesmen_workbook(make_workbook) -> Workbook:
    """Create an in-memory workbook with a canonical Salesmen sheet header."""

    return make_workbook(
        "Salesmen",
        ["SalesmanID", "SalesmanName", "IsActive"],
    )


@pytest.fixture
def transactions_workbook(make_workbook) -> Workbook:
    """Create an in-memory workbook with a canonical TransactionLog sheet header."""

    return make_workbook(
        "TransactionLog",
        [
            "TransactionID",
            "Timestamp",
            "TransactionType",
            "ProductID",
            "SalesmanID",
            "PaymentType",
            "QuantityChange",
            "TotalRevenue",
            "TotalCost",
            "LinkedTransactionID",
            "Notes",
        ],
    )


@pytest.fixture
def tmp_workbook_path(tmp_path: Path) -> Path:
    """Persist a minimal workbook on disk and return its path for file IO tests."""

    path = tmp_path / "master_workbook.xlsx"
    wb = openpyxl.Workbook()
    wb.save(path)
    return path
